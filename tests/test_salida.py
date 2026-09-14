"""Tests del escritor de Excel. Sinteticos: corren sin datos reales."""
import pathlib
import sys
import tempfile
import unittest

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / "src"))

from cotizador import salida  # noqa: E402
from cotizador.cartera import (agregar_para_salida, aplicar_escenario,  # noqa: E402
                               expandir_planes)
from cotizador.catalogos import GRUPOS_ETARIOS  # noqa: E402

TABLA = {g: {"optimista": 0.01, "pesimista": 0.02} for g in GRUPOS_ETARIOS}


def _bloques():
    base = expandir_planes([("unica", 30, "AMBA", None, 100.0),
                            ("unica", 1, "CORDOBA", None, 0.0)])
    return agregar_para_salida(aplicar_escenario(base, TABLA, "optimista"))


class TestEscritor(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        self.destino = pathlib.Path(self.dir.name) / "salida.xlsx"
        self.addCleanup(self.dir.cleanup)

    def _escribir(self, **kw):
        import openpyxl
        salida.escribir(self.destino, bloques_optimista=_bloques(),
                        bloques_pesimista=_bloques(), **kw)
        return openpyxl.load_workbook(self.destino, data_only=True)

    def test_las_ocho_hojas_en_orden(self):
        wb = self._escribir()
        self.assertEqual(wb.sheetnames, list(salida.HOJAS))
        self.assertEqual(len(wb.sheetnames), 8)
        wb.close()

    def test_las_columnas_de_cartera_son_las_esperadas(self):
        wb = self._escribir()
        ws = wb["Cartera optimista"]
        filas = list(ws.iter_rows(values_only=True))
        h = next(i for i, r in enumerate(filas) if r[0] == "subpoblacion")
        self.assertEqual(tuple(filas[h][:7]), salida.COLUMNAS_CARTERA)
        wb.close()

    def test_no_escribe_combinaciones_con_cantidad_cero(self):
        wb = self._escribir()
        ws = wb["Cartera optimista"]
        filas = list(ws.iter_rows(values_only=True))
        h = next(i for i, r in enumerate(filas) if r[0] == "subpoblacion")
        datos = [r for r in filas[h + 1:] if r[0]]
        self.assertTrue(datos)
        self.assertFalse([r for r in datos if r[6] == 0])
        wb.close()

    def test_cada_alternativa_cierra_en_su_poblacion(self):
        wb = self._escribir()
        ws = wb["Cartera optimista"]
        filas = list(ws.iter_rows(values_only=True))
        h = next(i for i, r in enumerate(filas) if r[0] == "subpoblacion")
        por_alt = {}
        for r in filas[h + 1:]:
            if not r[0]:
                continue
            por_alt[r[1]] = por_alt.get(r[1], 0.0) + r[6]
        self.assertEqual(len(por_alt), 9)
        for alt, v in por_alt.items():
            self.assertAlmostEqual(v, 100.0, delta=1e-9, msg=alt)
        wb.close()

    def test_la_hoja_avisa_que_no_se_suman_alternativas(self):
        wb = self._escribir()
        ws = wb["Cartera optimista"]
        primera = str(list(ws.iter_rows(values_only=True))[0][0])
        self.assertIn("NO SUMAR", primera)
        wb.close()

    def test_la_cartera_recibida_se_emite_aunque_este_vacia(self):
        wb = self._escribir(bloques_recibida={}, nota_recibida="sin situ observado")
        self.assertIn("Cartera recibida", wb.sheetnames)
        primera = str(list(wb["Cartera recibida"].iter_rows(values_only=True))[0][0])
        self.assertIn("sin situ observado", primera)
        wb.close()


if __name__ == "__main__":
    unittest.main()
