"""Escritor del Excel de salida. Ocho hojas de nombre y orden fijos.

Referencia: docs/handoff_v0.md, seccion 7.

Las hojas de cartera usan las cinco columnas y las denominaciones textuales
exactas de Referencia/referencia_tabla_a_completar.xlsx, mas `subpoblacion` y
`alternativa` al frente. Se omiten las combinaciones con cantidad cero.

Las nueve alternativas de plan se emiten siempre explicitas. Sus cantidades se
repiten cuando ninguna fuente informa plan: esa repeticion es intencional.
Las alternativas NO se suman entre si; el total de la licitacion es la suma de
las subpoblaciones dentro de una alternativa.
"""
from __future__ import annotations

import pathlib

HOJAS = (
    "Informacion recibida",
    "Distribuciones recibidas",
    "Diagnostico",
    "Supuestos",
    "Coberturas",
    "Cartera recibida",
    "Cartera optimista",
    "Cartera pesimista",
)

COLUMNAS_CARTERA = ("subpoblacion", "alternativa", "edad", "provincia",
                    "plan", "SITU", "cantidad")

_NOTA_ALTERNATIVAS = (
    "Las alternativas de plan son escenarios de cotizacion, no poblaciones "
    "distintas. NO SUMAR entre alternativas: el total de la licitacion es la "
    "suma de las subpoblaciones dentro de UNA alternativa."
)


def _escribir_tabla(ws, encabezado, filas, titulo=None):
    if titulo:
        ws.append([titulo])
        ws.append([])
    ws.append(list(encabezado))
    for fila in filas:
        ws.append(list(fila))


def _hoja_cartera(wb, nombre, bloques, nota=None):
    ws = wb.create_sheet(nombre)
    if nota:
        ws.append([nota])
        ws.append([])
    ws.append([_NOTA_ALTERNATIVAS])
    ws.append([])
    ws.append(list(COLUMNAS_CARTERA))
    n = 0
    for (sub, alt) in sorted(bloques):
        for clave in sorted(bloques[(sub, alt)]):
            grupo, provincia, plan, situ = clave
            cantidad = bloques[(sub, alt)][clave]
            if cantidad == 0:
                continue
            ws.append([sub, alt, grupo, provincia, plan, situ, cantidad])
            n += 1
    return n


def escribir(destino, *, bloques_optimista, bloques_pesimista,
             bloques_recibida=None, informacion_recibida=(),
             distribuciones=(), diagnostico=(), supuestos=(), coberturas=(),
             nota_recibida=None) -> dict:
    """Escribe el Excel de la licitacion. Devuelve el conteo de filas por hoja."""
    import openpyxl

    destino = pathlib.Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    conteo = {}

    ws = wb.create_sheet(HOJAS[0])
    enc, filas = (informacion_recibida[0], informacion_recibida[1]) \
        if informacion_recibida else ((), ())
    _escribir_tabla(ws, enc, filas)
    conteo[HOJAS[0]] = len(filas)

    ws = wb.create_sheet(HOJAS[1])
    enc, filas = (distribuciones[0], distribuciones[1]) if distribuciones else ((), ())
    _escribir_tabla(ws, enc, filas)
    conteo[HOJAS[1]] = len(filas)

    for nombre, contenido in ((HOJAS[2], diagnostico), (HOJAS[3], supuestos),
                              (HOJAS[4], coberturas)):
        ws = wb.create_sheet(nombre)
        enc, filas = (contenido[0], contenido[1]) if contenido else ((), ())
        _escribir_tabla(ws, enc, filas)
        conteo[nombre] = len(filas)

    conteo[HOJAS[5]] = _hoja_cartera(wb, HOJAS[5], bloques_recibida or {},
                                     nota=nota_recibida)
    conteo[HOJAS[6]] = _hoja_cartera(wb, HOJAS[6], bloques_optimista)
    conteo[HOJAS[7]] = _hoja_cartera(wb, HOJAS[7], bloques_pesimista)

    if list(wb.sheetnames) != list(HOJAS):
        raise RuntimeError(f"orden de hojas inesperado: {wb.sheetnames}")

    wb.save(destino)
    wb.close()
    return conteo
